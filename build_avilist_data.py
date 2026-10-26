#! /usr/bin/env python
"""Build src/AviList_data.csv — IUCN status, Birds of the World and IUCN links.

This is a BUILD-TIME script, not part of the app.  Run it by hand when AviList
publishes a new version.  It needs one package the app itself does not (install
it into a throwaway venv, not Yearbirder's):

    python3 -m venv /tmp/ybbuild && /tmp/ybbuild/bin/pip install openpyxl
    /tmp/ybbuild/bin/python build_avilist_data.py AviList-v2026-...-extended-version.xlsx

Source is the AviList EXTENDED checklist (.xlsx, https://www.avilist.org); the
short version lacks the columns used here.  AviList is CC BY 4.0, so the app
credits it wherever the data appears.

Output has one row per AviList species that carries a Cornell Lab species code
(the eBird SPECIES_CODE, which is how the app joins it):

    eBirdCode,Status,BoWURL,IUCNID

Status is AviList's IUCN Red List category as written there (LC, NT, VU, EN,
CR, "CR (PE)", "CR (PEW)", EW, EX, DD), or empty when it has none.  IUCNID is
the number at the end of AviList's BirdLife DataZone URL: BirdLife is the Red
List authority for birds, and its species ids are IUCN's taxon ids, so
https://www.iucnredlist.org/species/<IUCNID> opens the species' Red List page.

After rebuilding from a new AviList version, update AVILIST_YEAR in
src/code_DataBase.py, which labels the status in the Individual window.
"""

import csv
import os
import re
import sys

HERE     = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(HERE, "src", "AviList_data.csv")

COLUMNS = {"rank": "RANK",
           "code": "SPECIES CODE CORNELL LAB",
           "status": "IUCN RED LIST CATEGORY",
           "url": "BIRDS OF THE WORLD URL",
           "datazone": "BIRDLIFE DATAZONE URL"}


def main():
    if len(sys.argv) < 2:
        sys.exit("usage: build_avilist_data.py <AviList extended .xlsx>")
    import openpyxl

    xlsxPath = sys.argv[1]
    wb = openpyxl.load_workbook(xlsxPath, read_only=True)
    sheet = wb[wb.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows)]
    missing = [c for c in COLUMNS.values() if c not in header]
    if missing:
        sys.exit("Not an AviList extended file; missing columns: " + ", ".join(missing))
    idx = {k: header.index(c) for k, c in COLUMNS.items()}

    def cell(row, key):
        i = idx[key]
        return str(row[i]).strip() if i < len(row) and row[i] is not None else ""

    out = []
    seen = set()
    for row in rows:
        if cell(row, "rank") != "species":
            continue
        code = cell(row, "code")
        if not code or code in seen:
            continue
        seen.add(code)
        idMatch = re.search(r"/(\d+)/?$", cell(row, "datazone"))
        out.append((code, cell(row, "status"), cell(row, "url"),
                    idMatch.group(1) if idMatch else ""))

    with open(OUT_PATH, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["eBirdCode", "Status", "BoWURL", "IUCNID"])
        w.writerows(out)

    statuses = {}
    for _, s, _, _ in out:
        statuses[s or "(none)"] = statuses.get(s or "(none)", 0) + 1
    print("Sheet:", wb.sheetnames[0])
    year = re.search(r"v(\d{4})", wb.sheetnames[0] + os.path.basename(xlsxPath))
    print(f"Species written: {len(out)}   with IUCN id: {sum(1 for r in out if r[3])}")
    print("Status counts:", ", ".join(f"{k} {v}" for k, v in sorted(statuses.items())))
    if year:
        print(f"Check AVILIST_YEAR in src/code_DataBase.py is \"{year.group(1)}\".")
    print("Wrote", OUT_PATH)


if __name__ == "__main__":
    main()
